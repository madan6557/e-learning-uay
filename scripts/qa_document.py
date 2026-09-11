from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "docs" / "Technical Design - E-Learning UAY.docx"


def main() -> None:
    document = Document(str(DOC))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    with zipfile.ZipFile(DOC) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    alt_text = len(re.findall(r"<wp:docPr[^>]+(?:title|descr)=", xml))
    placeholders = {value: xml.count(value) for value in ("TODO", "TBD", "codex-file-citation")}
    print(f"document={DOC}")
    print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} inline_shapes={len(document.inline_shapes)}")
    print(f"headings={len(headings)} first={headings[:5]}")
    print(f"alt_text_attributes={alt_text}")
    print(f"placeholders={placeholders}")
    if len(document.inline_shapes) != 8 or alt_text != 8 or any(placeholders.values()):
        raise SystemExit("document QA failed")


if __name__ == "__main__":
    main()

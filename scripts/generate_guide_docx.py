"""Build the user book from the same role guide displayed by E-Learning UAY.
Use the Codex bundled Python (python-docx), then render with render_docx.py.
"""
from pathlib import Path
import argparse, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
LABELS = {'STUDENT':'Mahasiswa','INSTRUCTOR':'Dosen','DEPARTMENT_ADMIN':'Admin Prodi','SUPER_ADMIN':'Super Admin','RECTOR':'Rektor'}

def fill(cell, color):
    node=OxmlElement('w:shd'); node.set(qn('w:fill'),color); cell._tc.get_or_add_tcPr().append(node)
def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for vals in rows:
        for c,txt in zip(t.add_row().cells,vals):c.text=str(txt)
    for ri,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            pr.append(b);m=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'110');e.set(qn('w:type'),'dxa');m.append(e)
            pr.append(m);fill(c,'173E61' if ri==0 else ('F1F5F9' if ri%2 else 'FFFFFF'))
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.12
                for r in p.runs:r.font.size=Pt(10.5);r.font.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')

def new_heading(doc,title):
    p=doc.add_heading(title,1);p.paragraph_format.page_break_before=True

def build(output):
    guide=json.loads((ROOT/'apps/web/src/data/helpGuide.json').read_text(encoding='utf-8'))
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85)
    sec.header_distance=sec.footer_distance=Inches(.35)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    st=doc.styles['Normal'];st.font.size=Pt(11.5);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.14
    for name,size in [('Title',26),('Heading 1',20),('Heading 2',15),('Heading 3',12.5)]:
        st=doc.styles[name];st.font.size=Pt(size);st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(8)
    p=sec.header.paragraphs[0];p.text='E-Learning UAY  |  Panduan penggunaan  |  Edisi 7 Oktober 2026';p.style='Normal'
    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run('E-Learning UAY  |  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    p=doc.add_paragraph('UNIVERSITAS ACHMAD YANI BANJARMASIN');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    logo=ROOT/'docs/images/uay-logo.png'
    if logo.exists():
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(logo),width=Inches(1.5))
    doc.add_paragraph('Buku Panduan Penggunaan E Learning UAY',style='Title')
    doc.add_paragraph('Petunjuk bagi mahasiswa dosen administrator dan rektor',style='Subtitle')
    doc.add_paragraph('Buku ini menjelaskan penggunaan E-Learning UAY sesuai peran akun, mulai dari masuk dan mengikuti kelas hingga mengelola pembelajaran serta membaca laporan akademik. Gunakan bagian yang sesuai dengan tugas Anda.')
    doc.add_paragraph('Edisi 7 Oktober 2026\nhttps://e-learning.uay.ac.id\nTahun akademik 2026/2027')
    doc.add_paragraph('Panduan pada menu Bantuan memakai isi yang sama dengan buku ini. Pilih Buka buku panduan (tab baru) untuk membaca bagian yang mengikuti peran akun Anda. Langkah pada panduan HTML berupa daftar bernomor tanpa kotak centang.')
    new_heading(doc,'Cara menggunakan buku ini')
    doc.add_paragraph('Buku mencakup lima peran. Peran ditetapkan pengelola akun kampus; pengguna tidak dapat mengganti kewenangan melalui panduan. Dalam aplikasi, bagian yang tidak sesuai dengan peran tidak ditampilkan.')
    table(doc,['Peran','Bagian yang digunakan'],[(LABELS[role],', '.join(a['id'] for a in guide['tutorials'] if role in a['roles'])) for role in LABELS],[1.55,5.25])
    doc.add_paragraph('Tanggal dan jam di aplikasi serta laporan rektor mengikuti zona waktu perangkat. Periksa label zona atau offset, misalnya GMT+8. Mengubah zona tidak mengubah waktu kejadian aslinya. Pastikan pengaturan perangkat benar sebelum memasukkan jadwal.')
    doc.add_paragraph('Nama, jumlah, dan tanggal pada gambar adalah contoh. Label Demo — data simulasi menandakan lingkungan contoh. Laporan terhubung menggunakan data e-learning yang tersedia pada waktu data terakhir.')
    doc.add_heading('Daftar tutorial',2)
    for a in guide['tutorials']:doc.add_paragraph(a['id']+'  '+a['title'])
    for a in guide['tutorials']:
        new_heading(doc,a['id']+' '+a['title'].replace('&','dan'))
        doc.add_paragraph('Peran: '+', '.join(LABELS[r] for r in a['roles']))
        doc.add_paragraph(a['summary']);doc.add_paragraph('Buka: '+a['location']);doc.add_paragraph('Sebelum mulai: '+a['preparation'])
        if a['id']=='A4':
            for note in a['notes']:doc.add_paragraph('Catatan: '+note)
        doc.add_heading('Langkah penggunaan',2)
        for i,step in enumerate(a['steps'],1):
            p=doc.add_paragraph(f'{i}. {step}');p.paragraph_format.left_indent=Inches(.18);p.paragraph_format.first_line_indent=Inches(-.18)
        doc.add_paragraph('Periksa hasil: '+a['result'])
        if a.get('figure'):
            from PIL import Image
            f=a['figure'];img=ROOT/'apps/web/public'/f['src'].lstrip('/');w,h=Image.open(img).size
            width=min(6.75,(2.8 if a['id']=='A4' else 3.15)*w/h)
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(img),width=Inches(width))
            p=doc.add_paragraph(f['caption']+'. Gambar menggunakan data contoh.');p.paragraph_format.space_after=Pt(10)
            for r in p.runs:r.font.size=Pt(9.5)
        doc.add_heading('Pilihan dan tombol',2)
        table(doc,['Pilihan','Contoh','Kegunaan'],[(c['label'],c['value'],c['effect']) for c in a['controls']],[1.55,1.7,3.55])
        if a['notes'] and a['id']!='A4':
            doc.add_heading('Catatan penggunaan',2)
            for n in a['notes']:doc.add_paragraph(n)
    new_heading(doc,'Konversi nilai')
    doc.add_paragraph('Gunakan skala pada kelas. Jumlah bobot kategori harus 100 persen. Nilai akhir yang sudah diterbitkan mempertahankan versi kebijakannya. Pengubahan skala default bukan penghitungan ulang nilai yang telah terbit.')
    # Canonical scale values exported by the application, not copied from an old slide.
    scales=json.loads((ROOT/'docs/guide-grade-scales.json').read_text(encoding='utf-8'))
    for scale in scales.values():
        doc.add_heading('Skala '+scale['version'],2)
        table(doc,['Skor minimum','Huruf mutu','Indeks mutu'],[(str(b['minScore']),b['letter'],str(b['point'])) for b in scale['bands']],[2.4,2.2,2.2])
    new_heading(doc,'Solusi kendala')
    for a in guide['faqs']:
        doc.add_heading(a['title'],2);doc.add_paragraph('Peran: '+', '.join(LABELS[r] for r in a['roles']));doc.add_paragraph(a['content'])
    doc.core_properties.title='Buku Panduan Penggunaan E Learning UAY';doc.core_properties.subject='Petunjuk lima peran sesuai aplikasi pada 7 Oktober 2026';doc.core_properties.author='Universitas Achmad Yani';doc.core_properties.comments='Sumber isi apps/web/src/data/helpGuide.json. Gambar memakai data contoh.'
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output);print(output)
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--output',type=Path,default=ROOT/'docs/Buku Panduan Penggunaan E-Learning UAY.docx');args=ap.parse_args();build(args.output)
